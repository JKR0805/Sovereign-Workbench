"""DOCX renderer. The one complete artifact path.

Implemented with ``python-docx``. It produces a real file with headings, a
findings table, a metadata block, a citations list and an explicit unsupported
claims section. Section O: "Open it in Word. Not a preview. The real file."

If ``python-docx`` is absent, :meth:`DocxRenderer.status` says so and
:meth:`render` raises :class:`~vajra.core.exceptions.InfrastructureUnavailable`.
It never writes a placeholder file.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from vajra.artifacts.models import ArtifactSpec, BlockKind, DocumentBlock, RendererStatus
from vajra.core.enums import ArtifactKind
from vajra.core.exceptions import ArtifactError, InfrastructureUnavailable

MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _load_docx() -> Any:
    try:
        import docx
    except ImportError as exc:
        raise InfrastructureUnavailable(
            "python-docx is not installed. Install the 'artifacts' extra to enable "
            "DOCX generation.",
            extra="artifacts",
        ) from exc
    return docx


class DocxRenderer:
    """Renders an :class:`ArtifactSpec` to a ``.docx`` file."""

    kind = ArtifactKind.DOCX
    mime = MIME

    def status(self) -> RendererStatus:
        try:
            _load_docx()
        except InfrastructureUnavailable as exc:
            return RendererStatus(kind=self.kind, available=False, detail=exc.detail)
        return RendererStatus(
            kind=self.kind, available=True, detail="python-docx is available"
        )

    def render(self, spec: ArtifactSpec, destination: Path) -> Path:
        """Write the document and return the path it was written to."""
        docx = _load_docx()
        document = docx.Document()

        document.add_heading(spec.title, level=0)
        if spec.subtitle:
            document.add_paragraph(spec.subtitle)

        if spec.metadata:
            table = document.add_table(rows=0, cols=2)
            table.style = "Light Grid Accent 1"
            for key, value in spec.metadata.items():
                row = table.add_row().cells
                row[0].text = str(key)
                row[1].text = str(value)
            document.add_paragraph()

        for block in spec.blocks:
            self._render_block(document, block)

        if spec.citations:
            document.add_heading("Sources", level=1)
            for index, citation in enumerate(spec.citations, start=1):
                document.add_paragraph(f"[C{index}] {citation}", style="List Number")

        if spec.unsupported_claims:
            # Section G: unverifiable claims are marked, never dropped silently.
            document.add_heading("Unsupported claims", level=1)
            document.add_paragraph(
                "The following statements could not be verified against a cited source "
                "and must be checked before this document is relied upon."
            )
            for claim in spec.unsupported_claims:
                document.add_paragraph(claim, style="List Bullet")

        destination.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(destination))
        return destination

    def _render_block(self, document: Any, block: DocumentBlock) -> None:
        if block.kind is BlockKind.HEADING:
            document.add_heading(block.text or "", level=max(1, min(block.level, 4)))
        elif block.kind is BlockKind.PARAGRAPH:
            document.add_paragraph(block.text or "")
        elif block.kind is BlockKind.BULLETS:
            for item in block.items:
                document.add_paragraph(item, style="List Bullet")
        elif block.kind is BlockKind.KEY_VALUES:
            for key, value in block.pairs.items():
                paragraph = document.add_paragraph()
                paragraph.add_run(f"{key}: ").bold = True
                paragraph.add_run(str(value))
        elif block.kind is BlockKind.TABLE:
            self._render_table(document, block)
        elif block.kind is BlockKind.PAGE_BREAK:
            document.add_page_break()

    @staticmethod
    def _render_table(document: Any, block: DocumentBlock) -> None:
        if not block.columns:
            raise ArtifactError("A table block must declare its columns")
        table = document.add_table(rows=1, cols=len(block.columns))
        table.style = "Light Grid Accent 1"
        header = table.rows[0].cells
        for index, column in enumerate(block.columns):
            header[index].text = str(column)
            for paragraph in header[index].paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        for row in block.rows:
            if len(row) != len(block.columns):
                raise ArtifactError(
                    f"Table row has {len(row)} cells but {len(block.columns)} columns "
                    "were declared"
                )
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = str(value)
